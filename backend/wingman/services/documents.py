"""Context documents: résumé, job description and notes.

Text is extracted locally, split into chunks and searched with BM25. Small,
dependency-free and good enough for a handful of personal documents; it can be
swapped for embeddings later behind the same ``search`` method.
"""

from __future__ import annotations

import io
import math
import re
import uuid
from collections import Counter
from typing import Literal

from ..db import Database
from .conversations import now

DocKind = Literal["resume", "job_description", "notes"]
MAX_BYTES = 10 * 1024 * 1024
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
PROFILE_CHARS = 1500  # always-included head of résumé / job description

_TOKEN = re.compile(r"[a-z0-9][a-z0-9+#.\-]*")
_STOP = set(
    "a an the and or but if of to in on at for with by from as is are was were be been being it this that "
    "these those i you he she we they me my your our their do does did have has had not no so than then too "
    "very can could would should will just about into over under what which who whom how why when where".split()
)


class UnsupportedDocument(ValueError):
    pass


def tokenize(text: str) -> list[str]:
    return [t for t in _TOKEN.findall(text.lower()) if t not in _STOP and len(t) > 1]


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise UnsupportedDocument("File is larger than 10 MB.")
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        import docx

        document = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
    elif name.endswith((".txt", ".md", ".markdown", ".rtf", ".csv")) or "." not in name:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument("Use a PDF, Word (.docx), text or Markdown file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise UnsupportedDocument("No text found. If it's a scanned PDF, paste the text instead.")
    return text


def chunk_text(text: str) -> list[str]:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, current = [], ""
    for p in paragraphs:
        if len(current) + len(p) + 1 > CHUNK_CHARS and current:
            chunks.append(current)
            current = current[-CHUNK_OVERLAP:] + "\n" + p
        else:
            current = f"{current}\n{p}" if current else p
        while len(current) > CHUNK_CHARS * 1.5:
            chunks.append(current[:CHUNK_CHARS])
            current = current[CHUNK_CHARS - CHUNK_OVERLAP :]
    if current:
        chunks.append(current)
    return chunks


class BM25:
    def __init__(self, docs: list[list[str]], k1: float = 1.4, b: float = 0.75):
        self.docs = docs
        self.k1, self.b = k1, b
        self.avgdl = (sum(len(d) for d in docs) / len(docs)) if docs else 0
        df: Counter[str] = Counter()
        for d in docs:
            df.update(set(d))
        n = len(docs)
        self.idf = {t: math.log(1 + (n - f + 0.5) / (f + 0.5)) for t, f in df.items()}
        self.tf = [Counter(d) for d in docs]

    def scores(self, query: list[str]) -> list[float]:
        out = []
        for tf, doc in zip(self.tf, self.docs, strict=True):
            s = 0.0
            for t in query:
                if t not in tf:
                    continue
                f = tf[t]
                s += (
                    self.idf.get(t, 0)
                    * f
                    * (self.k1 + 1)
                    / (f + self.k1 * (1 - self.b + self.b * len(doc) / (self.avgdl or 1)))
                )
            out.append(s)
        return out


class DocumentStore:
    def __init__(self, db: Database):
        self.db = db
        self._index: tuple[list[dict], BM25] | None = None

    async def add(self, kind: DocKind, name: str, text: str) -> dict:
        doc_id = uuid.uuid4().hex[:12]
        chunks = chunk_text(text)
        await self.db.execute(
            "INSERT INTO documents(id, kind, name, chars, active, created_at) VALUES (?,?,?,?,1,?)",
            (doc_id, kind, name, len(text), now()),
        )
        await self.db.executemany(
            "INSERT INTO chunks(document_id, idx, text) VALUES (?,?,?)",
            [(doc_id, i, c) for i, c in enumerate(chunks)],
        )
        self._index = None
        return {"id": doc_id, "kind": kind, "name": name, "chars": len(text), "chunks": len(chunks), "active": True}

    async def list(self) -> list[dict]:
        rows = await self.db.fetchall(
            "SELECT d.*, (SELECT COUNT(*) FROM chunks c WHERE c.document_id = d.id) AS chunks"
            " FROM documents d ORDER BY d.created_at DESC"
        )
        for r in rows:
            r["active"] = bool(r["active"])
        return rows

    async def set_active(self, doc_id: str, active: bool) -> bool:
        changed = await self.db.execute("UPDATE documents SET active = ? WHERE id = ?", (int(active), doc_id))
        self._index = None
        return bool(changed)

    async def delete(self, doc_id: str) -> bool:
        deleted = await self.db.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        self._index = None
        return bool(deleted)

    async def _load_index(self) -> tuple[list[dict], BM25]:
        if self._index is None:
            rows = await self.db.fetchall(
                "SELECT c.text, c.idx, d.name, d.kind FROM chunks c JOIN documents d ON d.id = c.document_id"
                " WHERE d.active = 1 ORDER BY d.created_at, c.idx"
            )
            self._index = (rows, BM25([tokenize(r["text"]) for r in rows]))
        return self._index

    async def search(self, query: str, k: int = 4) -> list[dict]:
        rows, bm25 = await self._load_index()
        q = tokenize(query)
        if not rows or not q:
            return []
        scored = sorted(zip(bm25.scores(q), rows, strict=True), key=lambda x: x[0], reverse=True)
        return [dict(r, score=round(s, 3)) for s, r in scored[:k] if s > 0]

    async def profile(self) -> str:
        """Head of each active résumé and job description - always useful context."""
        rows = await self.db.fetchall(
            "SELECT d.kind, d.name, c.text FROM documents d JOIN chunks c ON c.document_id = d.id"
            " WHERE d.active = 1 AND d.kind IN ('resume', 'job_description') ORDER BY d.created_at, c.idx"
        )
        per_doc: dict[str, list[str]] = {}
        labels: dict[str, str] = {}
        for r in rows:
            key = r["name"]
            labels[key] = "Résumé" if r["kind"] == "resume" else "Job description"
            if sum(len(t) for t in per_doc.get(key, [])) < PROFILE_CHARS:
                per_doc.setdefault(key, []).append(r["text"])
        return "\n\n".join(f"[{labels[k]}: {k}]\n" + "\n".join(v)[:PROFILE_CHARS] for k, v in per_doc.items())

    async def context_for(self, question: str) -> str:
        profile = await self.profile()
        hits = await self.search(question, k=3)
        seen = profile
        extra = [f"[{h['name']}]\n{h['text']}" for h in hits if h["text"][:200] not in seen]
        return "\n\n".join(p for p in [profile, *extra] if p)
