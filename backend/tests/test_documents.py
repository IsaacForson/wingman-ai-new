import io

import docx
import pytest

from wingman.services.documents import DocumentStore, UnsupportedDocument, chunk_text, extract_text


def test_extracts_docx_and_text():
    d = docx.Document()
    d.add_paragraph("Senior engineer at Acme")
    d.add_paragraph("Built a Kafka pipeline")
    buf = io.BytesIO()
    d.save(buf)
    assert "Kafka pipeline" in extract_text("cv.docx", buf.getvalue())
    assert extract_text("notes.md", b"# Notes\n\nhello") == "# Notes\n\nhello"


def test_rejects_unknown_and_empty():
    with pytest.raises(UnsupportedDocument):
        extract_text("image.png", b"\x89PNG")
    with pytest.raises(UnsupportedDocument):
        extract_text("empty.txt", b"   ")


def test_chunks_long_text_with_overlap():
    text = "\n".join(f"Paragraph {i} " + "word " * 60 for i in range(30))
    chunks = chunk_text(text)
    assert len(chunks) > 5
    assert all(len(c) <= 1400 for c in chunks)


async def test_search_and_profile(db):
    store = DocumentStore(db)
    await store.add("resume", "cv.pdf", "Jane Doe\nSenior backend engineer. Scaled Kafka to 2M events/s at Acme.")
    await store.add("notes", "prep.md", "Company values: ownership and frugality.\nAsk about the on-call rotation.")
    hits = await store.search("tell me about kafka scaling")
    assert hits and hits[0]["name"] == "cv.pdf"
    profile = await store.profile()
    assert "Résumé: cv.pdf" in profile and "on-call" not in profile
    context = await store.context_for("what are the company values")
    assert "frugality" in context and "Jane Doe" in context


async def test_inactive_documents_are_ignored(db):
    store = DocumentStore(db)
    doc = await store.add("notes", "secret.md", "kubernetes migration plan")
    assert await store.search("kubernetes")
    await store.set_active(doc["id"], False)
    assert await store.search("kubernetes") == []
    assert await store.delete(doc["id"])
    assert await store.list() == []
